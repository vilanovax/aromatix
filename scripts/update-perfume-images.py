#!/usr/bin/env python3
"""Set perfume.document.image from the image sheet (status=دارد → 800 URL).

Slugs in the sheet without an approved photo lose document.image.
"""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from urllib.parse import unquote, urlparse

ROOT = Path(__file__).resolve().parents[1]
def default_xlsx() -> Path:
    downloads = Path.home() / "Downloads"
    for name in ("perfume-images.v2.xlsx", "perfume-images.v1.xlsx"):
        path = downloads / name
        if path.is_file():
            return path
    return downloads / "perfume-images.v2.xlsx"


XLSX = Path(sys.argv[1] if len(sys.argv) > 1 else default_xlsx())


def load_image_sheet(path: Path) -> tuple[dict[str, str], set[str]]:
    """Approved 800 URLs, and every slug listed in the sheet."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["عکس عطرها"]
    approved: dict[str, str] = {}
    listed: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        slug = row[1]
        status = row[5]
        link_800 = row[8]
        if not slug:
            continue
        listed.add(str(slug))
        if status != "دارد" or not link_800:
            continue
        url = str(link_800).strip()
        if url.startswith("https://"):
            approved[str(slug)] = url
    wb.close()
    return approved, listed


def db_connect():
    import psycopg2

    line = next(
        l for l in (ROOT / ".env").read_text().splitlines() if l.startswith("DATABASE_URL=")
    )
    p = urlparse(line.split("=", 1)[1].strip().strip('"').strip("'"))
    return psycopg2.connect(
        host=p.hostname,
        port=p.port,
        user=unquote(p.username or ""),
        password=unquote(p.password or ""),
        dbname=p.path.lstrip("/"),
        connect_timeout=30,
    )


def main() -> None:
    if not XLSX.is_file():
        raise SystemExit(f"missing xlsx: {XLSX}")

    images, listed = load_image_sheet(XLSX)
    version = "v2" if "v2" in XLSX.name else "v1"
    map_out = ROOT / "data" / "catalog" / f"perfume-images.{version}.json"
    map_out.parent.mkdir(parents=True, exist_ok=True)
    map_out.write_text(
        json.dumps(
            {
                "_meta": {
                    "source": str(XLSX.name),
                    "status": "دارد",
                    "field": "image",
                    "url": "https://c977199.parspack.net/perfumes/800/{slug}.webp",
                    "count": len(images),
                    "listed": len(listed),
                },
                "images": images,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"map written {map_out} n={len(images)} listed={len(listed)}", flush=True)

    t0 = time.time()
    conn = db_connect()
    print(f"connected {time.time() - t0:.1f}s", flush=True)
    cur = conn.cursor()

    slugs = sorted(images)
    # Only touch rows that exist and have a storage URL in the sheet.
    cur.execute(
        'SELECT slug FROM aroma."CatalogPerfume" WHERE slug = ANY(%s)',
        (slugs,),
    )
    in_db = {r[0] for r in cur.fetchall()}
    missing_in_db = sorted(set(slugs) - in_db)
    targets = [s for s in slugs if s in in_db]
    print(f"targets_in_db={len(targets)} xlsx_not_in_db={len(missing_in_db)}", flush=True)

    updated = 0
    already = 0
    batch_size = 50
    for i in range(0, len(targets), batch_size):
        chunk = targets[i : i + batch_size]
        cur.execute(
            'SELECT slug, document FROM aroma."CatalogPerfume" WHERE slug = ANY(%s)',
            (chunk,),
        )
        for slug, document in cur.fetchall():
            url = images[slug]
            try:
                doc = json.loads(document)
            except json.JSONDecodeError:
                print(f"bad json: {slug}", flush=True)
                continue
            if doc.get("image") == url:
                already += 1
                continue
            doc["image"] = url
            cur.execute(
                'UPDATE aroma."CatalogPerfume" SET document = %s, "updatedAt" = NOW() WHERE slug = %s',
                (json.dumps(doc, ensure_ascii=False), slug),
            )
            updated += 1
        conn.commit()
        print(f"batch {i // batch_size + 1}: updated_total={updated}", flush=True)

    # Drop storage URLs for slugs the sheet no longer approves.
    cleared = 0
    cur.execute(
        """
        SELECT slug, document FROM aroma."CatalogPerfume"
        WHERE document LIKE %s
        """,
        ('%"image":%',),
    )
    for slug, document in cur.fetchall():
        if slug not in listed or slug in images:
            continue
        try:
            doc = json.loads(document)
        except json.JSONDecodeError:
            continue
        if "image" not in doc:
            continue
        doc.pop("image", None)
        cur.execute(
            'UPDATE aroma."CatalogPerfume" SET document = %s, "updatedAt" = NOW() WHERE slug = %s',
            (json.dumps(doc, ensure_ascii=False), slug),
        )
        cleared += 1
    conn.commit()
    print(f"cleared_unapproved={cleared}", flush=True)

    cur.execute(
        """
        SELECT count(*) FROM aroma."CatalogPerfume"
        WHERE document LIKE %s
        """,
        ('%"image": "https://c977199.parspack.net/perfumes/800/%',),
    )
    with_image = cur.fetchone()[0]
    conn.close()

    print(
        f"updated={updated} already={already} cleared={cleared} db_with_storage_image={with_image} "
        f"xlsx_not_in_db={len(missing_in_db)}",
        flush=True,
    )
    if missing_in_db[:8]:
        print("sample not in db:", ", ".join(missing_in_db[:8]), flush=True)


if __name__ == "__main__":
    main()
